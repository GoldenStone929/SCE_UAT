from datetime import datetime
from pathlib import Path
import csv
import html
import importlib.util
import json
import sys


def get_root():
    if len(sys.argv) >= 2:
        return Path(sys.argv[1]).resolve()
    candidate = Path(__file__).resolve().parents[2]
    if (candidate / "data").exists():
        return candidate
    return Path.cwd().resolve()


def add_result(rows, output_type, file_path, status, package_used="", error_message=""):
    rows.append(
        {
            "language": "Python",
            "output_type": output_type,
            "file_path": str(file_path) if file_path else "",
            "status": status,
            "package_used": package_used,
            "error_message": error_message,
        }
    )


def file_ok(path):
    return path.exists() and path.stat().st_size > 0


def optional_available(module_name):
    return importlib.util.find_spec(module_name) is not None


def main():
    root = get_root()
    output_dir = root / "outputs" / "generated_documents"
    result_file = root / "outputs" / "test_results" / "python_document_generation_results.csv"
    output_dir.mkdir(parents=True, exist_ok=True)
    result_file.parent.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().isoformat(timespec="seconds")
    rows = []

    try:
        path = output_dir / "python_document_test.txt"
        path.write_text(f"Python TXT document generation PASS\nTimestamp: {timestamp}\n", encoding="utf-8")
        add_result(rows, "TXT", path, "PASS" if file_ok(path) else "FAIL", "standard_library")
    except Exception as exc:
        add_result(rows, "TXT", "", "ERROR", "standard_library", str(exc))

    try:
        path = output_dir / "python_document_test.csv"
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(["item", "value"])
            writer.writerow(["timestamp", timestamp])
            writer.writerow(["status", "PASS"])
        add_result(rows, "CSV", path, "PASS" if file_ok(path) else "FAIL", "standard_library")
    except Exception as exc:
        add_result(rows, "CSV", "", "ERROR", "standard_library", str(exc))

    try:
        path = output_dir / "python_document_test.json"
        path.write_text(
            json.dumps({"timestamp": timestamp, "status": "PASS"}, indent=2),
            encoding="utf-8",
        )
        add_result(rows, "JSON", path, "PASS" if file_ok(path) else "FAIL", "standard_library")
    except Exception as exc:
        add_result(rows, "JSON", "", "ERROR", "standard_library", str(exc))

    try:
        path = output_dir / "python_document_test.html"
        path.write_text(
            "<!doctype html><html><head><meta charset='utf-8'><title>Python Document Test</title></head>"
            f"<body><h1>Python Document Test</h1><p>{html.escape(timestamp)}</p></body></html>",
            encoding="utf-8",
        )
        add_result(rows, "HTML", path, "PASS" if file_ok(path) else "FAIL", "standard_library")
    except Exception as exc:
        add_result(rows, "HTML", "", "ERROR", "standard_library", str(exc))

    if optional_available("openpyxl"):
        try:
            from openpyxl import Workbook

            path = output_dir / "python_document_test.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "UAT"
            sheet.append(["item", "value"])
            sheet.append(["timestamp", timestamp])
            sheet.append(["status", "PASS"])
            workbook.save(path)
            add_result(rows, "XLSX", path, "PASS" if file_ok(path) else "FAIL", "openpyxl")
        except Exception as exc:
            add_result(rows, "XLSX", "", "ERROR", "openpyxl", str(exc))
    else:
        add_result(rows, "XLSX", "", "NOT_AVAILABLE", "openpyxl")

    if optional_available("docx"):
        try:
            from docx import Document

            path = output_dir / "python_document_test.docx"
            document = Document()
            document.add_heading("Python Document Test", level=1)
            document.add_paragraph(f"Timestamp: {timestamp}")
            document.add_paragraph("Status: PASS")
            document.save(path)
            add_result(rows, "DOCX", path, "PASS" if file_ok(path) else "FAIL", "docx")
        except Exception as exc:
            add_result(rows, "DOCX", "", "ERROR", "docx", str(exc))
    else:
        add_result(rows, "DOCX", "", "NOT_AVAILABLE", "docx")

    if optional_available("reportlab"):
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen import canvas

            path = output_dir / "python_document_test.pdf"
            pdf = canvas.Canvas(str(path), pagesize=letter)
            pdf.drawString(72, 720, "Python Document Test")
            pdf.drawString(72, 700, f"Timestamp: {timestamp}")
            pdf.drawString(72, 680, "Status: PASS")
            pdf.save()
            add_result(rows, "PDF", path, "PASS" if file_ok(path) else "FAIL", "reportlab")
        except Exception as exc:
            add_result(rows, "PDF", "", "ERROR", "reportlab", str(exc))
    else:
        add_result(rows, "PDF", "", "NOT_AVAILABLE", "reportlab")

    with result_file.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["language", "output_type", "file_path", "status", "package_used", "error_message"],
        )
        writer.writeheader()
        writer.writerows(rows)

    baseline_statuses = [row["status"] for row in rows if row["output_type"] in {"TXT", "CSV", "JSON", "HTML"}]
    print(f"Python document generation results written to {result_file}")
    return 0 if all(status == "PASS" for status in baseline_statuses) else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"Python document generation failed: {exc}", file=sys.stderr)
        raise SystemExit(1)
